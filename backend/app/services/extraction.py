from pathlib import Path


def extract_text(filename: str, payload: bytes) -> tuple[str | None, list[str]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return payload.decode("utf-8", errors="replace"), []
    try:
        import io
        if suffix == ".pdf":
            from pypdf import PdfReader
            return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(payload)).pages), []
        if suffix == ".docx":
            from docx import Document
            return "\n".join(p.text for p in Document(io.BytesIO(payload)).paragraphs), []
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            return "\n".join(" ".join(str(cell.value or "") for cell in row) for sheet in workbook for row in sheet.iter_rows()), []
        if suffix == ".pptx":
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(payload))
            return "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")), []
    except ImportError:
        return None, [f"Optional extractor for {suffix} is not installed"]
    except Exception as exc:
        return None, [f"Could not extract {suffix}: {exc.__class__.__name__}"]
    return None, ["No text extractor is available for this file"]
